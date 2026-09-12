import streamlit as st
import pypdf
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import re
from collections import defaultdict
from datetime import datetime
import pandas as pd
import io

st.set_page_config(page_title="GST Universal Converter", page_icon="📑", layout="wide")

st.markdown("<h2 style='text-align: center; color: #1F4E79;'>📑 Xenium Tyres - GST Universal Invoices to Excel & Google Sheets</h2>", unsafe_allow_html=True)
st.markdown("<p style='text-align: center; color: #555;'>दोनों PDF फॉर्मेट सपोर्टेड (Multiprint व Smart GST) | किसी बाहरी टेम्पलेट की आवश्यकता नहीं</p>", unsafe_allow_html=True)
st.write("---")

uploaded_pdf = st.file_uploader("📂 अपना इनवॉइस PDF चुनें (Single या Multiple Pages)", type=["pdf"])

def parse_invoices(file_bytes):
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    invoices = []
    
    for page_idx, page in enumerate(reader.pages):
        text = page.extract_text()
        if not text or "TAX INVOICE" not in text:
            continue
            
        # 1. Invoice Number
        inv_no_m = re.search(r'Invoice\s*No\.?\s*[:\|\s]?\s*(\d+)', text)
        if not inv_no_m:
            inv_no_m = re.search(r'Invoice\s*No\.?\s*\n\s*(\d+)', text)
        inv_no = int(inv_no_m.group(1)) if inv_no_m else (page_idx + 1)
        
        # 2. Invoice Date
        inv_date_m = re.search(r'Invoice\s*Date\s*[:\|\s]?\s*([0-9]{1,2}[\.\-\/][0-9A-Za-z]{2,3}[\.\-\/][0-9]{2,4})', text)
        if not inv_date_m:
            inv_date_m = re.search(r'Invoice\s*Date\s*\n\s*([0-9]{1,2}[\.\-\/][0-9A-Za-z]{2,3}[\.\-\/][0-9]{2,4})', text)
        if not inv_date_m:
            inv_date_m = re.search(r'(?:Date:?|ORIGINAL FOR RECIPIENT)\s*\n?\s*([0-9]{1,2}[\.\-\/][0-9A-Za-z]{2,3}[\.\-\/][0-9]{2,4})', text)
            
        inv_date = ""
        if inv_date_m:
            raw_d = inv_date_m.group(1).strip()
            try:
                if '.' in raw_d:
                    dt = datetime.strptime(raw_d, "%d.%m.%Y")
                    inv_date = dt.strftime("%d-%b-%Y")
                elif '-' in raw_d:
                    parts = raw_d.split('-')
                    if len(parts[1]) == 2:
                        dt = datetime.strptime(raw_d, "%d-%m-%Y")
                        inv_date = dt.strftime("%d-%b-%Y")
                    else:
                        inv_date = raw_d
                else:
                    inv_date = raw_d
            except:
                inv_date = raw_d
                
        # 3. Customer Details & GSTIN
        cust_name = "Unregistered Consumer"
        cust_addr = ""
        cust_gstin = ""
        
        cust_blk = re.search(r'Customer Detail\s*\n(.*?)(?:Place of\s*Supply|TAX INVOICE|Invoice No|Due Date)', text, re.DOTALL)
        if cust_blk:
            cb = cust_blk.group(1)
            gst_all = re.findall(r'[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z]{1}[1-9A-Z]{1}Z[0-9A-Z]{1}', cb)
            for g in gst_all:
                if g != "09KQTPS5380E1ZW":
                    cust_gstin = g
                    break
                    
            nm = re.search(r'(?:Name|M\/S|MS\/)\s*[:\|\s]?\s*\n?\s*([^\n]+)', cb, re.IGNORECASE)
            if nm:
                c_cand = nm.group(1).strip()
                if not any(k in c_cand.lower() for k in ["address", "phone", "gstin", "customer"]):
                    cust_name = c_cand
                    
            ad = re.search(r'Address\s*[:\|\s]?\s*\n?\s*([^\n]+(?:\n\s*[^\n]+)?)', cb, re.IGNORECASE)
            if ad:
                lines = [l.strip() for l in ad.group(1).split('\n') if l.strip() and not any(k in l.lower() for k in ['phone', 'gstin', 'place of'])]
                cust_addr = " ".join(lines)
                
        # 4. Totals
        taxable_m = re.search(r'Taxable Amount\s*[:\|\s]?\s*\n?\s*([\d,\.]+)', text)
        cgst_m = re.search(r'Add\s*:\s*CGST\s*[:\|\s]?\s*\n?\s*([\d,\.]+)', text)
        sgst_m = re.search(r'Add\s*:\s*SGST\s*[:\|\s]?\s*\n?\s*([\d,\.]+)', text)
        tot_m = re.search(r'Total Amount After Tax\s*[:\|\s]?\s*\n?\s*[₹\s]*([\d,\.]+)', text)
        
        taxable_val = float(taxable_m.group(1).replace(',', '')) if taxable_m else 0.0
        cgst_val = float(cgst_m.group(1).replace(',', '')) if cgst_m else 0.0
        sgst_val = float(sgst_m.group(1).replace(',', '')) if sgst_m else 0.0
        tot_val = float(tot_m.group(1).replace(',', '')) if tot_m else (taxable_val + cgst_val + sgst_val)
        
        # 5. Items Extraction
        items_area_m = re.search(r'Sr\.\s*\n?No\..*?(?:Total\s*\n?\s*[\d\.]+\s*(?:PCS|NOS))', text, re.DOTALL)
        items = []
        if items_area_m:
            area = items_area_m.group(0)
            hsn_matches = list(re.finditer(r'(40\d{6})\s*\n?\s*(\d+(?:\.\d+)?)\s*(PCS|NOS)?\s*\n?\s*([\d,\.]+)\s*\n?\s*([\d,\.]+)\s*\n?\s*([\d,\.]+)\s*\n?\s*([\d,\.]+)\s*\n?\s*([\d,\.]+)\s*\n?\s*([\d,\.]+)\s*\n?\s*([\d,\.]+)', area))
            for idx_h, hm in enumerate(hsn_matches):
                hsn, qty, uqc_unit, rate, tax_amt, cgst_r, cgst_a, sgst_r, sgst_a, item_tot = hm.groups()
                
                h_start = hm.start()
                start_search = hsn_matches[idx_h - 1].end() if idx_h > 0 else 0
                desc_text = area[start_search:h_start]
                d_lines = [l.strip() for l in desc_text.split('\n') if l.strip()]
                clean_words = []
                for l in d_lines:
                    if l.isdigit() and len(l) <= 2: continue
                    if any(bad in l for bad in ["Xenium", "Tyres", "Ahraura", "Mirzapur", "231301", "Sr.", "No.", "Name of Product", "HSN", "Qty", "Rate", "Taxable Value", "CGST", "SGST", "Total", "% Amount"]): continue
                    clean_words.append(l)
                item_desc = " ".join(clean_words).strip() or f"Tyre / Tube (HSN {hsn})"
                uqc = "NOS-NUMBERS" if (uqc_unit and "NOS" in uqc_unit) else "PCS-PIECES"
                
                items.append({
                    'desc': item_desc,
                    'hsn': int(hsn),
                    'qty': float(qty),
                    'uqc': uqc,
                    'rate': float(rate.replace(',', '')),
                    'taxable': float(tax_amt.replace(',', '')),
                    'cgst_r': float(cgst_r.replace(',', '')),
                    'cgst_a': float(cgst_a.replace(',', '')),
                    'sgst_r': float(sgst_r.replace(',', '')),
                    'sgst_a': float(sgst_a.replace(',', '')),
                    'total': float(item_tot.replace(',', ''))
                })
                
        invoices.append({
            'invoice_no': inv_no,
            'invoice_date': inv_date,
            'customer_name': cust_name,
            'address': cust_addr,
            'cust_gstin': cust_gstin,
            'pos': "09-Uttar Pradesh",
            'taxable_val': taxable_val,
            'cgst_val': cgst_val,
            'sgst_val': sgst_val,
            'tot_val': tot_val,
            'items': items
        })
        
    return invoices

def create_full_gstr1_excel(invoices):
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet
    
    b2b_invs = [x for x in invoices if x['cust_gstin']]
    b2c_invs = [x for x in invoices if not x['cust_gstin']]
    
    # 1. Sheet: b2b,sez,de
    ws_b2b = wb.create_sheet('b2b,sez,de')
    ws_b2b.append(['Summary For B2B(4)'] + [None]*11 + ['HELP'])
    ws_b2b.append(['No. of Recipients', None, 'No. of Invoices', None, 'Total Invoice Value', None, None, None, None, None, None, 'Total Taxable Value', 'Total Cess'])
    ws_b2b.append([
        len(set(x['cust_gstin'] for x in b2b_invs)), None,
        len(b2b_invs), None,
        round(sum(x['tot_val'] for x in b2b_invs), 2), None, None, None, None, None, None,
        round(sum(x['taxable_val'] for x in b2b_invs), 2), 0.0
    ])
    ws_b2b.append(['GSTIN/UIN of Recipient', 'Receiver Name', 'Invoice Number', 'Invoice date', 'Invoice Value', 'Place Of Supply', 'Reverse Charge', 'Applicable % of Tax Rate', 'Invoice Type', 'E-Commerce GSTIN', 'Rate', 'Taxable Value', 'Cess Amount'])
    for inv in b2b_invs:
        ws_b2b.append([inv['cust_gstin'], inv['customer_name'], inv['invoice_no'], inv['invoice_date'], inv['tot_val'], inv['pos'], 'N', None, 'Regular B2B', None, 18.0, inv['taxable_val'], 0.0])

    # 2. Sheet: b2cs
    ws_b2cs = wb.create_sheet('b2cs')
    ws_b2cs.append(['Summary For B2CS(7)'] + [None]*5 + ['HELP'])
    ws_b2cs.append([None, None, None, None, 'Total Taxable  Value', 'Total Cess', None])
    tot_b2c_tax = round(sum(x['taxable_val'] for x in b2c_invs), 2)
    ws_b2cs.append([None, None, None, None, tot_b2c_tax, 0.0, None])
    ws_b2cs.append(['Type', 'Place Of Supply', 'Applicable % of Tax Rate', 'Rate', 'Taxable Value', 'Cess Amount', 'E-Commerce GSTIN'])
    if tot_b2c_tax > 0:
        ws_b2cs.append(['OE', '09-Uttar Pradesh', None, 18.0, tot_b2c_tax, 0.0, None])

    # 3. Sheet: hsn(b2b)
    ws_hb = wb.create_sheet('hsn(b2b)')
    ws_hb.append(['Summary For HSN(12)'] + [None]*8 + ['HELP', None])
    ws_hb.append(['No. of HSN', None, None, None, 'Total Value', None, 'Total Taxable Value', 'Total Integrated Tax', 'Total Central Tax', 'Total State/UT Tax', 'Total Cess'])
    hsn_b2b = defaultdict(lambda: {'desc': set(), 'uqc': 'PCS-PIECES', 'qty': 0.0, 'total': 0.0, 'taxable': 0.0, 'cgst': 0.0, 'sgst': 0.0})
    for inv in b2b_invs:
        for it in inv['items']:
            h = it['hsn']
            hsn_b2b[h]['desc'].add(it['desc'])
            hsn_b2b[h]['qty'] += it['qty']
            hsn_b2b[h]['uqc'] = it['uqc']
            hsn_b2b[h]['total'] += it['total']
            hsn_b2b[h]['taxable'] += it['taxable']
            hsn_b2b[h]['cgst'] += it['cgst_a']
            hsn_b2b[h]['sgst'] += it['sgst_a']
    ws_hb.append([
        len(hsn_b2b), None, None, None,
        round(sum(v['total'] for v in hsn_b2b.values()), 2), None,
        round(sum(v['taxable'] for v in hsn_b2b.values()), 2), 0.0,
        round(sum(v['cgst'] for v in hsn_b2b.values()), 2),
        round(sum(v['sgst'] for v in hsn_b2b.values()), 2), 0.0
    ])
    ws_hb.append(['HSN', 'Description', 'UQC', 'Total Quantity', 'Total Value', 'Rate', 'Taxable Value', 'Integrated Tax Amount', 'Central Tax Amount', 'State/UT Tax Amount', 'Cess Amount'])
    for hsn, data in sorted(hsn_b2b.items()):
        ws_hb.append([int(hsn), ", ".join(sorted(data['desc'])), data['uqc'], data['qty'], round(data['total'], 2), 18.0, round(data['taxable'], 2), 0.0, round(data['cgst'], 2), round(data['sgst'], 2), 0.0])

    # 4. Sheet: hsn(b2c)
    ws_hc = wb.create_sheet('hsn(b2c)')
    ws_hc.append(['Summary For HSN(12)'] + [None]*8 + ['HELP', None])
    ws_hc.append(['No. of HSN', None, None, None, 'Total Value', None, 'Total Taxable Value', 'Total Integrated Tax', 'Total Central Tax', 'Total State/UT Tax', 'Total Cess'])
    hsn_b2c = defaultdict(lambda: {'desc': set(), 'uqc': 'PCS-PIECES', 'qty': 0.0, 'total': 0.0, 'taxable': 0.0, 'cgst': 0.0, 'sgst': 0.0})
    for inv in b2c_invs:
        for it in inv['items']:
            h = it['hsn']
            hsn_b2c[h]['desc'].add(it['desc'])
            hsn_b2c[h]['qty'] += it['qty']
            hsn_b2c[h]['uqc'] = it['uqc']
            hsn_b2c[h]['total'] += it['total']
            hsn_b2c[h]['taxable'] += it['taxable']
            hsn_b2c[h]['cgst'] += it['cgst_a']
            hsn_b2c[h]['sgst'] += it['sgst_a']
    ws_hc.append([
        len(hsn_b2c), None, None, None,
        round(sum(v['total'] for v in hsn_b2c.values()), 2), None,
        round(sum(v['taxable'] for v in hsn_b2c.values()), 2), 0.0,
        round(sum(v['cgst'] for v in hsn_b2c.values()), 2),
        round(sum(v['sgst'] for v in hsn_b2c.values()), 2), 0.0
    ])
    ws_hc.append(['HSN', 'Description', 'UQC', 'Total Quantity', 'Total Value', 'Rate', 'Taxable Value', 'Integrated Tax Amount', 'Central Tax Amount', 'State/UT Tax Amount', 'Cess Amount'])
    for hsn, data in sorted(hsn_b2c.items()):
        ws_hc.append([int(hsn), ", ".join(sorted(data['desc'])), data['uqc'], data['qty'], round(data['total'], 2), 18.0, round(data['taxable'], 2), 0.0, round(data['cgst'], 2), round(data['sgst'], 2), 0.0])

    # 5. Sheet: docs
    ws_docs = wb.create_sheet('docs')
    ws_docs.append(['Summary of documents issued during the tax period (13)', None, None, None, 'HELP'])
    ws_docs.append([None, None, None, 'Total Number', 'Total Cancelled'])
    inv_nums = [x['invoice_no'] for x in invoices if x['invoice_no']]
    ws_docs.append([None, None, None, len(inv_nums), 0])
    ws_docs.append(['Nature of Document', 'Sr. No. From', 'Sr. No. To', 'Total Number', 'Cancelled'])
    ws_docs.append(['Invoices for outward supply', min(inv_nums) if inv_nums else 1, max(inv_nums) if inv_nums else len(inv_nums), len(inv_nums), 0])

    # 6. Sheet: Invoice_Register (पूर्ण बिल विवरण)
    ws_reg = wb.create_sheet('Invoice_Register')
    headers = ['Invoice No', 'Date', 'Customer Name', 'Address', 'GSTIN', 'Item Desc', 'HSN', 'Qty', 'Unit', 'Rate', 'Taxable Value', 'CGST (9%)', 'SGST (9%)', 'Total Invoice']
    ws_reg.append(headers)
    for col_idx in range(1, len(headers)+1):
        c = ws_reg.cell(1, col_idx)
        c.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        c.font = Font(name="Calibri", bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center")
        
    for inv in invoices:
        for it in inv['items']:
            ws_reg.append([
                inv['invoice_no'], inv['invoice_date'], inv['customer_name'],
                inv['address'], inv['cust_gstin'] or 'Unregistered', it['desc'],
                it['hsn'], it['qty'], it['uqc'], it['rate'], it['taxable'],
                it['cgst_a'], it['sgst_a'], it['total']
            ])

    return wb

if uploaded_pdf and st.button("🚀 Convert to GSTR-1 Excel & Sheets", type="primary", use_container_width=True):
    with st.spinner("PDF को पूरी तरह स्कैन व एक्सट्रैक्ट किया जा रहा है..."):
        invoices = parse_invoices(uploaded_pdf.read())
        
        if not invoices:
            st.error("कोई बिल नहीं पढ़ा जा सका। कृपया वैध इनवॉइस PDF चुनें।")
        else:
            wb = create_full_gstr1_excel(invoices)
            
            excel_buffer = io.BytesIO()
            wb.save(excel_buffer)
            excel_buffer.seek(0)
            
            # DataFrame for Google Sheets
            flat_data = []
            for inv in invoices:
                for it in inv['items']:
                    flat_data.append({
                        'Invoice No': inv['invoice_no'],
                        'Date': inv['invoice_date'],
                        'Customer Name': inv['customer_name'],
                        'Address': inv['address'],
                        'GSTIN': inv['cust_gstin'] or 'Unregistered',
                        'Item': it['desc'],
                        'HSN': it['hsn'],
                        'Qty': it['qty'],
                        'Rate': it['rate'],
                        'Taxable Value': it['taxable'],
                        'CGST': it['cgst_a'],
                        'SGST': it['sgst_a'],
                        'Total Amount': it['total']
                    })
            df = pd.DataFrame(flat_data)
            csv_buffer = df.to_csv(index=False).encode('utf-8')

            b2b_c = sum(1 for x in invoices if x['cust_gstin'])
            b2c_c = sum(1 for x in invoices if not x['cust_gstin'])
            total_sales = sum(x['tot_val'] for x in invoices)

            st.success(f"✅ सफलता! कुल {len(invoices)} बिल एक्सट्रैक्ट हुए (B2B: {b2b_c}, B2C: {b2c_c}) | कुल बिक्री: ₹{total_sales:,.2f}")
            
            # Download & Google Sheets Row
            st.subheader("📥 डेटा डाउनलोड व Google Sheets विकल्प")
            col1, col2, col3 = st.columns(3)
            with col1:
                st.download_button(
                    label="📊 Download GSTR-1 Excel (.xlsx)",
                    data=excel_buffer,
                    file_name="GSTR1_Filled_Ready.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary",
                    use_container_width=True
                )
            with col2:
                st.download_button(
                    label="📄 Download for Google Sheets (.csv)",
                    data=csv_buffer,
                    file_name="Invoices_For_GoogleSheets.csv",
                    mime="text/csv",
                    use_container_width=True
                )
            with col3:
                st.link_button(
                    label="🌐 Open Google Sheets",
                    url="https://sheets.new",
                    use_container_width=True
                )

            st.markdown("---")
            st.subheader("👁️ लाइव डेटा प्रीव्यू (Google Sheets की तरह)")
            st.dataframe(df, use_container_width=True)
            
            st.info("💡 **Google Sheets में खोलने का 1-क्लिक तरीका:**\n1. ऊपर 'Download for Google Sheets (.csv)' या Excel पर क्लिक करें।\n2. 'Open Google Sheets' बटन दबाएँ।\n3. Google Sheets में **File ➔ Import ➔ Upload** से फ़ाइल चुनें; सारा डेटा तुरंत शीट में आ जाएगा!")
